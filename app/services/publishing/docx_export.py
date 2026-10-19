"""Word manuscript (DOCX): accepted by KDP for ebook and paperback conversion, and handy for human editors."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

from .manuscript import Manuscript, inline_segments


def _runs(paragraph, text: str) -> None:
    for segment, italic in inline_segments(text):
        run = paragraph.add_run(segment)
        run.italic = italic


def _centered(doc, text: str, size: int, *, bold=False, italic=False, space_before=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.first_line_indent = Inches(0)
    run = p.add_run(text)
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    return p


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_docx(ms: Manuscript, output: Path) -> dict:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.first_line_indent = Inches(0.3)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.15
    pf.widow_control = True
    for section in doc.sections:
        section.page_width, section.page_height = Inches(6), Inches(9)
        section.left_margin = section.right_margin = Inches(0.8)
        section.top_margin = section.bottom_margin = Inches(0.8)

    core = doc.core_properties
    core.title, core.author = ms.title, ms.author
    core.subject = ms.subtitle or ""
    core.keywords = ", ".join(ms.keywords)

    _centered(doc, ms.title, 28, bold=True, space_before=160)
    if ms.subtitle:
        _centered(doc, ms.subtitle, 14, italic=True, space_before=8)
    _centered(doc, ms.author, 16, space_before=120)
    _page_break(doc)
    for line in ms.copyright_lines:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(line)
        r.font.size = Pt(9)
    if ms.dedication:
        _page_break(doc)
        _centered(doc, ms.dedication, 12, italic=True, space_before=160)

    for part in ms.parts:
        doc.add_section(WD_SECTION.NEW_PAGE)
        if part.label:
            _centered(doc, part.label.upper(), 11, space_before=90)
        if part.title:
            h = doc.add_heading(level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            h.paragraph_format.space_after = Pt(30)
            h.add_run(part.title)
        for b in part.blocks:
            if b.kind == "break":
                _centered(doc, "* * *", 12, space_before=10)
                continue
            p = doc.add_paragraph()
            if b.kind == "verse":
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.space_after = Pt(12)
                lines = b.text.split("\n")
                for i, line in enumerate(lines):
                    _runs(p, line)
                    if i < len(lines) - 1:
                        p.add_run().add_break()
            else:
                if b.first:
                    p.paragraph_format.first_line_indent = Inches(0)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _runs(p, b.text)

    for heading, text in (("Acknowledgments", ms.acknowledgments), ("About the Author", ms.author_bio)):
        if text:
            doc.add_section(WD_SECTION.NEW_PAGE)
            h = doc.add_heading(heading, level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for para in [t for t in text.split("\n\n") if t.strip()]:
                _runs(doc.add_paragraph(), para)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return {"valid": True}


def build_txt(ms: Manuscript, output: Path) -> dict:
    from .manuscript import to_plain

    lines = [ms.title.upper()]
    if ms.subtitle:
        lines.append(ms.subtitle)
    lines += [f"by {ms.author}", "", *ms.copyright_lines, ""]
    if ms.dedication:
        lines += [ms.dedication, ""]
    for part in ms.parts:
        lines += ["", "", " ".join(x for x in (part.label, part.title) if x).upper() if part.label else part.title.upper(), ""]
        for b in part.blocks:
            lines += ["* * *" if b.kind == "break" else to_plain(b.text), ""]
    for heading, text in (("ACKNOWLEDGMENTS", ms.acknowledgments), ("ABOUT THE AUTHOR", ms.author_bio)):
        if text:
            lines += ["", heading, "", text, ""]
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return {"valid": True}
